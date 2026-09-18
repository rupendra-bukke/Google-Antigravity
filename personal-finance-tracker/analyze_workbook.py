#!/usr/bin/env python3
"""
Analyze an existing personal-finance Excel workbook and suggest structural improvements.

Usage:
    python analyze_workbook.py /path/to/your/file.xlsx
"""

from __future__ import annotations

import argparse
import re
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

RECOMMENDED_COLUMNS = {
    "loans": {
        "required": ["Entry_Date", "Loan_Name", "Lender", "EMI_Amount", "Amount_Paid", "Status"],
        "optional": ["Principal", "Interest_Rate", "Outstanding_Balance", "Payment_Method", "Notes"],
    },
    "school": {
        "required": ["Entry_Date", "Child_Name", "School", "Fee_Type", "Amount_Due", "Amount_Paid", "Status"],
        "optional": ["Academic_Year", "Term", "Due_Date", "Payment_Date", "Payment_Method", "Notes"],
    },
    "emi": {
        "required": ["Entry_Date", "Item_Name", "EMI_Amount", "Status"],
        "optional": ["Category", "Vendor_Lender", "Tenure_Months", "EMIs_Remaining", "Next_Due_Date", "Notes"],
    },
    "other": {
        "required": ["Entry_Date", "Category", "Description", "Amount", "Status"],
        "optional": ["Subcategory", "Payment_Method", "Recurring", "Notes"],
    },
}


def normalize_header(value: object) -> str:
    if value is None:
        return ""
    text = str(value).strip().lower()
    text = re.sub(r"[^a-z0-9]+", "_", text)
    return text.strip("_")


def guess_sheet_type(sheet_name: str, headers: list[str]) -> str:
    name = sheet_name.lower()
    header_text = " ".join(headers)
    if any(k in name for k in ("loan", "mortgage", "home")) or "emi_amount" in header_text:
        return "loans"
    if any(k in name for k in ("school", "fee", "tuition", "kid", "child")):
        return "school"
    if "emi" in name or "installment" in name:
        return "emi"
    return "other"


def find_header_row(ws, max_scan: int = 20) -> tuple[int, list[str]]:
    best_row = 1
    best_headers: list[str] = []
    best_score = 0
    for row in range(1, max_scan + 1):
        values = [ws.cell(row=row, column=col).value for col in range(1, ws.max_column + 1)]
        non_empty = [v for v in values if v not in (None, "")]
        if len(non_empty) >= best_score and len(non_empty) >= 3:
            best_score = len(non_empty)
            best_row = row
            best_headers = [normalize_header(v) for v in values if v not in (None, "")]
    return best_row, best_headers


def analyze_sheet(ws) -> dict:
    header_row, headers = find_header_row(ws)
    data_rows = 0
    numeric_cols: dict[int, int] = defaultdict(int)
    date_cols: dict[int, int] = defaultdict(int)
    empty_rows = 0

    for row in range(header_row + 1, ws.max_row + 1):
        row_vals = [ws.cell(row=row, column=col).value for col in range(1, len(headers) + 1)]
        if all(v in (None, "") for v in row_vals):
            empty_rows += 1
            continue
        data_rows += 1
        for col, val in enumerate(row_vals, start=1):
            if isinstance(val, (int, float)):
                numeric_cols[col] += 1
            if hasattr(val, "year"):
                date_cols[col] += 1

    sheet_type = guess_sheet_type(ws.title, headers)
    spec = RECOMMENDED_COLUMNS[sheet_type]
    missing_required = [h for h in spec["required"] if h.lower() not in headers]
    missing_optional = [h for h in spec["optional"] if h.lower() not in headers]

    issues = []
    if data_rows == 0:
        issues.append("No data rows detected below the header.")
    if missing_required:
        issues.append(f"Missing recommended columns: {', '.join(missing_required)}")
    if len(headers) > 20:
        issues.append("Too many columns — consider splitting summary fields into a separate sheet.")
    if empty_rows > data_rows and data_rows > 0:
        issues.append("Many blank rows between data — use an Excel Table and filter blanks.")

    mixed_summary = any(k in " ".join(headers) for k in ("total", "balance", "summary")) and data_rows > 5
    if mixed_summary:
        issues.append("Sheet may mix transaction rows with summary rows — keep one row per transaction.")

    return {
        "sheet": ws.title,
        "type": sheet_type,
        "header_row": header_row,
        "headers": headers,
        "data_rows": data_rows,
        "missing_required": missing_required,
        "missing_optional": missing_optional,
        "issues": issues,
        "has_dashboard": "dashboard" in ws.title.lower() or "summary" in ws.title.lower(),
    }


def print_report(results: list[dict], path: Path) -> None:
    print("=" * 72)
    print(f"Workbook analysis: {path.name}")
    print("=" * 72)

    for r in results:
        print(f"\n## Sheet: {r['sheet']}  (detected type: {r['type']})")
        print(f"   Header row: {r['header_row']}  |  Data rows: {r['data_rows']}")
        print(f"   Columns found: {', '.join(r['headers']) or '(none)'}")
        if r["missing_required"]:
            print(f"   [!] Add columns: {', '.join(r['missing_required'])}")
        if r["missing_optional"]:
            print(f"   [i] Nice to have: {', '.join(r['missing_optional'])}")
        if r["issues"]:
            for issue in r["issues"]:
                print(f"   [!] {issue}")
        if not r["has_dashboard"]:
            print(f"   [i] Suggested dashboard sheet: {r['sheet']}_Dashboard")

    print("\n" + "=" * 72)
    print("GENERAL RECOMMENDATIONS")
    print("=" * 72)
    print(
        """
1. One row = one transaction/payment (never put monthly totals in the same table as line items).
2. Separate DATA sheets from DASHBOARD sheets (formulas + charts reference the data table).
3. Use consistent column names across sheets: Entry_Date, Amount_Paid, Status, Payment_Method, Notes.
4. Convert each data range into an Excel Table (Insert > Table) so formulas auto-expand.
5. Add a Setup sheet for dropdown lists (payment methods, statuses, categories).
6. Use Status values: Paid, Pending, Overdue, Partial — enables pending-item counts on dashboards.
7. Keep an Overview_Dashboard that sums all categories for a monthly family finance snapshot.

To migrate: run create_finance_template.py, then copy your data into the matching sheets.
"""
    )


def main() -> None:
    parser = argparse.ArgumentParser(description="Analyze a personal finance Excel workbook.")
    parser.add_argument("workbook", type=Path, help="Path to .xlsx file")
    args = parser.parse_args()

    if not args.workbook.exists():
        raise SystemExit(f"File not found: {args.workbook}")

    wb = load_workbook(args.workbook, data_only=True)
    results = [analyze_sheet(wb[sheet]) for sheet in wb.sheetnames]
    print_report(results, args.workbook.resolve())


if __name__ == "__main__":
    main()
